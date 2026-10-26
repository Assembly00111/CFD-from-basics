#!/usr/bin/env python3
"""
generate_workbooks.py
---------------------
Builds three 2D CFD / field-solver workbooks that run entirely on Excel
formulas (no macros):

    1_heat_conduction_2d.xlsx   transient 2D heat conduction  (explicit FTCS)
    2_potential_flow_2d.xlsx    2D potential flow, stream function (Gauss-Seidel)
    3_lid_driven_cavity_2d.xlsx 2D incompressible Navier-Stokes (psi-omega)

The "time loop" is Excel's iterative calculation: every cell updates from
itself and its neighbours, and each Excel iteration is one solver step.

Usage:
    python generate_workbooks.py                       # defaults
    python generate_workbooks.py --out build --n-cavity 41 --n-heat 31

Requires: openpyxl >= 3.1
"""
import argparse
import math
import os

from openpyxl import Workbook
from openpyxl.chart import Reference, ScatterChart, Series
from openpyxl.formatting.rule import ColorScaleRule, FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter as L
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.workbook.properties import CalcProperties

FONT = "Arial"
F_BASE = Font(name=FONT, size=10)
F_BOLD = Font(name=FONT, size=10, bold=True)
F_TITLE = Font(name=FONT, size=14, bold=True)
F_INPUT = Font(name=FONT, size=10, color="0000FF")
F_GRID = Font(name=FONT, size=8)
F_HDR = Font(name=FONT, size=8, bold=True)
FILL_INPUT = PatternFill("solid", start_color="FFFF99")
FILL_HDR = PatternFill("solid", start_color="D9D9D9")
FILL_DARK = PatternFill("solid", start_color="1F1F1F", end_color="1F1F1F")
FILL_GREY = PatternFill("solid", start_color="808080", end_color="808080")

GRID_R0 = 3  # first grid row   (row 2 holds x-axis, column A holds y-axis)
GRID_C0 = 2  # first grid column (column B)


# --------------------------------------------------------------------------
# generic helpers
# --------------------------------------------------------------------------
def new_workbook():
    wb = Workbook()
    wb.remove(wb.active)
    # Iterative calculation = the solver engine.
    wb.calculation = CalcProperties(
        iterate=True, iterateCount=1000, iterateDelta=1e-9, fullCalcOnLoad=True
    )
    return wb


def cell(i, j):
    """A1 address of grid node (row i from the top, column j from the left)."""
    return f"{L(GRID_C0 + j)}{GRID_R0 + i}"


def grid_range(NX, NY):
    return f"{cell(0, 0)}:{cell(NY - 1, NX - 1)}"


def make_inputs_sheet(wb, title, rows, notes=None):
    """rows: list of (label, name, value, note, kind); kind in input|calc|state|section.
    Creates defined names p_<name> pointing at column B."""
    ws = wb.create_sheet("Inputs")
    ws["A1"] = title
    ws["A1"].font = F_TITLE
    ws["A2"] = "Blue on yellow = you may edit.  Black = calculated by the workbook."
    ws["A2"].font = Font(name=FONT, size=9, italic=True)
    r = 4
    for label, name, value, note, kind in rows:
        if kind == "section":
            ws.cell(r, 1, label).font = F_BOLD
            ws.cell(r, 1).fill = FILL_HDR
            ws.cell(r, 2).fill = FILL_HDR
            ws.cell(r, 3).fill = FILL_HDR
            r += 1
            continue
        ws.cell(r, 1, label).font = F_BASE
        c = ws.cell(r, 2, value)
        c.font = F_INPUT if kind == "input" else F_BASE
        if kind == "input":
            c.fill = FILL_INPUT
        ws.cell(r, 3, note).font = Font(name=FONT, size=9, color="595959")
        if name:
            wb.defined_names[f"p_{name}"] = DefinedName(
                f"p_{name}", attr_text=f"Inputs!$B${r}"
            )
        r += 1
    ws.column_dimensions["A"].width = 38
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 90
    return ws


def make_instructions(wb, title, lines):
    ws = wb.create_sheet("Instructions", 0)
    ws["A1"] = title
    ws["A1"].font = F_TITLE
    for k, line in enumerate(lines, start=3):
        c = ws.cell(k, 1, line)
        c.font = F_BOLD if line.isupper() or line.endswith(":") else F_BASE
        c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 120
    return ws


def make_field_sheet(wb, name, NX, NY, fmt, title):
    """Sheet with x header (row 2), y header (col A) and a styled empty grid."""
    ws = wb.create_sheet(name)
    ws["A1"] = title
    ws["A1"].font = F_BOLD
    ws["A2"] = "y \\ x"
    ws["A2"].font = F_HDR
    ws["A2"].fill = FILL_HDR
    for j in range(NX):
        c = ws.cell(2, GRID_C0 + j)
        c.value = 0 if j == 0 else f"={L(GRID_C0 + j - 1)}2+p_h"
        c.number_format = "0.00"
        c.font = F_HDR
        c.fill = FILL_HDR
        c.alignment = Alignment(horizontal="center")
        ws.column_dimensions[L(GRID_C0 + j)].width = 7.5
    for i in range(NY):
        c = ws.cell(GRID_R0 + i, 1)
        c.value = f"={NY - 1}*p_h" if i == 0 else f"=A{GRID_R0 + i - 1}-p_h"
        c.number_format = "0.00"
        c.font = F_HDR
        c.fill = FILL_HDR
        for j in range(NX):
            g = ws.cell(GRID_R0 + i, GRID_C0 + j)
            g.number_format = fmt
            g.font = F_GRID
            g.alignment = Alignment(horizontal="center")
    ws.column_dimensions["A"].width = 7
    ws.freeze_panes = "B3"
    return ws


def add_colorscale(ws, rng, diverging=False, mid=0):
    if diverging:
        rule = ColorScaleRule(
            start_type="min", start_color="2C7BB6",
            mid_type="num", mid_value=mid, mid_color="FFFFFF",
            end_type="max", end_color="D7191C",
        )
    else:
        rule = ColorScaleRule(
            start_type="min", start_color="2C7BB6",
            mid_type="percentile", mid_value=50, mid_color="FFFFBF",
            end_type="max", end_color="D7191C",
        )
    ws.conditional_formatting.add(rng, rule)


def add_streamline_bands(ws, NX, NY, name_spacing="p_dpsi", width=0.12):
    """Dark bands wherever psi is within `width` of a multiple of dspsi -> streamlines."""
    rng = grid_range(NX, NY)
    first = cell(0, 0)
    ws.conditional_formatting.add(
        rng,
        FormulaRule(
            formula=[f"MOD({first}/{name_spacing},1)<{width}"],
            fill=FILL_DARK,
            font=Font(name=FONT, size=8, color="FFFFFF"),
        ),
    )


def monitor_chart(ws, anchor, title, x_title, y_title, series, width=14, height=9):
    """series: list of (ws, x_col, y_col, r0, r1, label, markers_only)"""
    ch = ScatterChart()
    ch.title = title
    ch.style = 13
    ch.x_axis.title = x_title
    ch.y_axis.title = y_title
    ch.width, ch.height = width, height
    ch.x_axis.delete = False
    ch.y_axis.delete = False
    for sws, xc, yc, r0, r1, label, markers_only in series:
        xs = Reference(sws, min_col=xc, min_row=r0, max_row=r1)
        ys = Reference(sws, min_col=yc, min_row=r0, max_row=r1)
        s = Series(ys, xs, title=label)
        if markers_only:
            s.marker.symbol = "circle"
            s.marker.size = 7
            s.graphicalProperties.line.noFill = True
        else:
            s.marker.symbol = "none"
            s.smooth = False
        ch.series.append(s)
    ws.add_chart(ch, anchor)


# --------------------------------------------------------------------------
# 1. transient 2D heat conduction
# --------------------------------------------------------------------------
def build_heat(path, N=21):
    NX = NY = N
    wb = new_workbook()
    rows = [
        ("RUN CONTROL", None, None, None, "section"),
        ("Reset (1 = initialise, 0 = run)", "reset", 0,
         "Set to 1 to restore the initial temperature, then back to 0 to run.", "input"),
        ("BOUNDARY / INITIAL CONDITIONS", None, None, None, "section"),
        ("T top edge [deg C]", "Ttop", 100, "Dirichlet", "input"),
        ("T bottom edge [deg C]", "Tbottom", 0, "Dirichlet", "input"),
        ("T left edge [deg C]", "Tleft", 0, "Dirichlet", "input"),
        ("T right edge [deg C]", "Tright", 0, "Dirichlet", "input"),
        ("Initial interior T [deg C]", "Tinit", 0, "Applied while Reset = 1", "input"),
        ("PHYSICS / NUMERICS", None, None, None, "section"),
        ("Thermal diffusivity alpha [m^2/s]", "alpha", 1.0, "Plate is 1 m x 1 m, so the diffusion time scale is ~1 s", "input"),
        ("Time step dt [s]", "dt", 0.0005, "Keep Fo <= 0.25 or the explicit scheme blows up", "input"),
        ("Grid nodes per side N", "N", N, "Fixed by the sheet layout; re-run generate_workbooks.py for another N", "calc"),
        ("Grid spacing h [m]", "h", "=1/(p_N-1)", "Plate side length = 1 m", "calc"),
        ("Fourier number Fo = alpha*dt/h^2", "Fo", "=p_alpha*p_dt/p_h^2", "Stability limit of FTCS in 2D: 0.25", "calc"),
        ("Stability check", None, '=IF(p_Fo<=0.25,"STABLE","UNSTABLE - reduce dt")', "", "calc"),
        ("MONITORS", None, None, None, "section"),
        ("Mean temperature [deg C]", None, f"=AVERAGE(T!{grid_range(NX, NY)})", "Plateaus when steady state is reached", "calc"),
        ("Max temperature [deg C]", None, f"=MAX(T!{grid_range(NX, NY)})", "", "calc"),
        ("Min temperature [deg C]", None, f"=MIN(T!{grid_range(NX, NY)})", "", "calc"),
        ("Max |error| vs exact centreline (y<1)", None, "=Validation!H3", "Only meaningful at steady state with 3 edges at 0 deg C (see Validation sheet)", "calc"),
    ]
    inp = make_inputs_sheet(wb, "2D transient heat conduction  -  inputs & monitors", rows)

    ws = make_field_sheet(wb, "T", NX, NY, "0.0", "Temperature field T [deg C]   (row 3 = top edge y=1, last row = bottom edge y=0)")
    for i in range(NY):
        for j in range(NX):
            a = cell(i, j)
            if i == 0 and j in (0, NX - 1):
                f = "=AVERAGE(p_Ttop,p_Tleft)" if j == 0 else "=AVERAGE(p_Ttop,p_Tright)"
            elif i == NY - 1 and j in (0, NX - 1):
                f = "=AVERAGE(p_Tbottom,p_Tleft)" if j == 0 else "=AVERAGE(p_Tbottom,p_Tright)"
            elif i == 0:
                f = "=p_Ttop"
            elif i == NY - 1:
                f = "=p_Tbottom"
            elif j == 0:
                f = "=p_Tleft"
            elif j == NX - 1:
                f = "=p_Tright"
            else:
                E, W, Nn, S = cell(i, j + 1), cell(i, j - 1), cell(i - 1, j), cell(i + 1, j)
                f = f"=IF(p_reset=1,p_Tinit,{a}+p_Fo*({E}+{W}+{Nn}+{S}-4*{a}))"
            ws[a] = f
    add_colorscale(ws, grid_range(NX, NY))

    # ---- validation against the exact Fourier-series solution (centre line x = 0.5)
    v = wb.create_sheet("Validation")
    v["A1"] = "Centre-line (x = 0.5) temperature vs exact series solution"
    v["A1"].font = F_BOLD
    hdr = ["y [m]", "T simulated", "T exact", "|error|"]
    for k, t in enumerate(hdr):
        c = v.cell(2, 1 + k, t)
        c.font = F_BOLD
        c.fill = FILL_HDR
    mid = NX // 2
    for i in range(NY):
        r = 3 + i
        v.cell(r, 1, f"=T!A{GRID_R0 + i}")
        v.cell(r, 2, f"=T!{L(GRID_C0 + mid)}{GRID_R0 + i}")
        v.cell(r, 3, f"=(4*p_Ttop/PI())*SUMPRODUCT(SIN($K$3:$K$42*PI()*0.5)*SINH($K$3:$K$42*PI()*A{r})/SINH($K$3:$K$42*PI())/$K$3:$K$42)")
        v.cell(r, 4, f"=ABS(B{r}-C{r})")
        for c in range(1, 5):
            v.cell(r, c).font = F_BASE
            v.cell(r, c).number_format = "0.000"
    v["G2"] = "Summary"
    v["G2"].font = F_BOLD
    v["G3"] = "Max |error| (y<1)"
    v["H3"] = f'=IF(AND(p_Tbottom=0,p_Tleft=0,p_Tright=0),MAX(D4:D{2 + NY}),"n/a")'
    v["G4"] = "Valid when"
    v["H4"] = "T bottom = T left = T right = 0, and the run has reached steady state"
    v["J2"] = "Odd n"
    v["K2"] = "n"
    v["J2"].font = F_BOLD
    v["K2"].font = F_BOLD
    for k in range(40):
        v.cell(3 + k, 11, 2 * k + 1).font = F_BASE
    v["J3"] = "Series terms"
    for a in ("G3", "G4", "H3", "H4", "J3"):
        v[a].font = F_BASE
    v["H3"].number_format = "0.000"
    v["G14"] = ("Exact: T(x,y) = (4*Ttop/pi) * sum_{n odd} sin(n*pi*x) * sinh(n*pi*y) / (sinh(n*pi) * n), "
                "evaluated at x = 0.5 with 40 terms. The top row (y = 1) is excluded from the max error because the series itself converges slowly at the heated edge.")
    v["G14"].font = Font(name=FONT, size=9, italic=True)
    v.column_dimensions["G"].width = 14
    v.column_dimensions["H"].width = 14
    monitor_chart(v, "G16", "Centre-line temperature", "T [deg C]", "y [m]",
                  [(v, 2, 1, 3, 2 + NY, "Spreadsheet", False),
                   (v, 3, 1, 3, 2 + NY, "Exact series", True)])

    make_instructions(wb, "2D transient heat conduction in Excel", [
        "WHAT THIS SOLVES",
        "dT/dt = alpha * (d2T/dx2 + d2T/dy2) on a 1 m x 1 m plate with fixed (Dirichlet) edge temperatures.",
        "Scheme: explicit FTCS.  T_new = T + Fo*(T_E + T_W + T_N + T_S - 4*T),  Fo = alpha*dt/h^2 <= 0.25.",
        "",
        "ONE-TIME SETUP (required):",
        "1) File > Options > Formulas > tick 'Enable iterative calculation'.  Set Maximum Iterations = 1000, Maximum Change = 0.000000001.",
        "   (Excel takes this setting from the first workbook opened in a session. Close other workbooks first if you see a circular-reference warning.)",
        "2) Press F9. Each F9 advances the solution by 'Maximum Iterations' time steps. Watch the 'T' sheet heat map evolve.",
        "",
        "HOW TO USE:",
        "- Change edge temperatures / alpha / dt on the 'Inputs' sheet, then press F9 repeatedly until the monitors stop changing.",
        "- Set Reset = 1 to wipe the field, then 0 to run again.",
        "- 'Validation' compares the centre line with the exact Fourier series (top edge hot, other three edges 0).",
        "",
        "TRY: set Tleft=100 and Tright=0 for a left-to-right gradient; raise dt until the stability check shows UNSTABLE and watch it explode.",
    ])
    wb.move_sheet("Inputs", offset=-(len(wb.sheetnames) - 2))
    wb.save(path)


# --------------------------------------------------------------------------
# 2. 2D potential flow (stream-function form) over a bump
# --------------------------------------------------------------------------
def build_potential(path, NX=41, NY=21):
    h = 1.0 / (NY - 1)
    wb = new_workbook()
    rows = [
        ("RUN CONTROL", None, None, None, "section"),
        ("Reset (1 = initialise, 0 = run)", "reset", 0,
         "Set to 1 to restore psi = U*y everywhere, then back to 0.", "input"),
        ("FLOW", None, None, None, "section"),
        ("Free-stream speed U [m/s]", "U", 1.0, "Uniform inlet flow in +x", "input"),
        ("Streamline spacing d(psi)", "dpsi", 0.05, "Dark bands on the Psi sheet are lines of constant psi", "input"),
        ("GRID", None, None, None, "section"),
        ("Nodes in x (NX)", "NX", NX, "Fixed by layout; re-run generate_workbooks.py to change", "calc"),
        ("Nodes in y (NY)", "NY", NY, "Channel height = 1 m, length = (NX-1)/(NY-1) m", "calc"),
        ("Grid spacing h [m]", "h", "=1/(p_NY-1)", "", "calc"),
        ("MONITORS", None, None, None, "section"),
        ("Max speed [m/s]", None, f"=MAX(Speed!{grid_range(NX, NY)})", "Ideal flow accelerates over the bump crest", "calc"),
        ("Min speed (fluid only) [m/s]", None, f'=MIN(Speed!{grid_range(NX, NY)})', "Low near the stagnation points ahead of / behind the bump", "calc"),
        ("Max |v| (vertical velocity) [m/s]", None, f"=MAX(MAX(v_vel!{grid_range(NX, NY)}),-MIN(v_vel!{grid_range(NX, NY)}))", "Vertical deflection of the stream; stops changing when converged", "calc"),
    ]
    make_inputs_sheet(wb, "2D potential flow over a bump  -  inputs & monitors", rows)

    # ---- mask sheet (user-editable obstacle)
    mk = make_field_sheet(wb, "Mask", NX, NY, "0", "Obstacle mask: 1 = solid, 0 = fluid.  Edit freely (solids must touch the bottom wall).")
    for i in range(NY):
        y = (NY - 1 - i) * h
        for j in range(NX):
            x = j * h
            solid = 1 if ((x - 1.0) ** 2 + y ** 2 <= 0.3 ** 2) and 0 < j < NX - 1 else 0
            c = mk[cell(i, j)]
            c.value = solid
            c.font = Font(name=FONT, size=8, color="0000FF")
    mk.conditional_formatting.add(grid_range(NX, NY), FormulaRule(formula=[f"{cell(0, 0)}=1"], fill=FILL_DARK, font=Font(name=FONT, size=8, color="FFFFFF")))

    # ---- psi
    ps = make_field_sheet(wb, "Psi", NX, NY, "0.000", "Stream function psi   (u = dpsi/dy, v = -dpsi/dx)")
    for i in range(NY):
        for j in range(NX):
            a = cell(i, j)
            ycell = f"$A{GRID_R0 + i}"
            if i == NY - 1:
                f = 0
            elif i == 0:
                f = f"=p_U*{ycell}"
            elif j == 0:
                f = f"=p_U*{ycell}"
            elif j == NX - 1:
                f = f"=IF(p_reset=1,p_U*{ycell},{cell(i, j - 1)})"  # zero-gradient outlet
            else:
                E, W, Nn, S = cell(i, j + 1), cell(i, j - 1), cell(i - 1, j), cell(i + 1, j)
                f = f"=IF(Mask!{a}=1,0,IF(p_reset=1,p_U*{ycell},0.25*({E}+{W}+{Nn}+{S})))"
            ps[a] = f
    add_streamline_bands(ps, NX, NY)
    add_colorscale(ps, grid_range(NX, NY))

    # ---- velocities
    def vel_formula(i, j, comp):
        a = cell(i, j)
        if comp == "u":
            if i == 0:
                core = f"(Psi!{cell(0, j)}-Psi!{cell(1, j)})/p_h"
            elif i == NY - 1:
                core = f"(Psi!{cell(NY - 2, j)}-Psi!{cell(NY - 1, j)})/p_h"
            else:
                core = f"(Psi!{cell(i - 1, j)}-Psi!{cell(i + 1, j)})/(2*p_h)"
        else:
            if j == 0:
                core = f"-(Psi!{cell(i, 1)}-Psi!{cell(i, 0)})/p_h"
            elif j == NX - 1:
                core = f"-(Psi!{cell(i, NX - 1)}-Psi!{cell(i, NX - 2)})/p_h"
            else:
                core = f"-(Psi!{cell(i, j + 1)}-Psi!{cell(i, j - 1)})/(2*p_h)"
        return f"=IF(Mask!{a}=1,0,{core})"

    uv = make_field_sheet(wb, "u_vel", NX, NY, "0.00", "Horizontal velocity u [m/s]")
    vv = make_field_sheet(wb, "v_vel", NX, NY, "0.00", "Vertical velocity v [m/s]")
    sp = make_field_sheet(wb, "Speed", NX, NY, "0.00", "Speed |V| [m/s]   (dark = solid)")
    for i in range(NY):
        for j in range(NX):
            a = cell(i, j)
            uv[a] = vel_formula(i, j, "u")
            vv[a] = vel_formula(i, j, "v")
            sp[a] = f'=IF(Mask!{a}=1,"",SQRT(u_vel!{a}^2+v_vel!{a}^2))'
    add_colorscale(uv, grid_range(NX, NY))
    add_colorscale(vv, grid_range(NX, NY), diverging=True)
    add_colorscale(sp, grid_range(NX, NY))
    sp.conditional_formatting.add(grid_range(NX, NY), FormulaRule(formula=[f'ISTEXT({cell(0, 0)})'], fill=FILL_GREY))

    make_instructions(wb, "2D potential flow in Excel (stream-function formulation)", [
        "WHAT THIS SOLVES",
        "Irrotational, incompressible flow:  laplacian(psi) = 0,  with u = dpsi/dy, v = -dpsi/dx.",
        "Scheme: Gauss-Seidel.  psi = 0.25*(psi_E + psi_W + psi_N + psi_S) at every fluid node.",
        "Boundaries: inlet psi = U*y;  bottom wall and solid bump psi = 0;  top wall psi = U*H;  outlet: zero gradient.",
        "",
        "ONE-TIME SETUP (required):",
        "1) File > Options > Formulas > tick 'Enable iterative calculation'.  Set Maximum Iterations = 1000, Maximum Change = 0.000000001.",
        "   (Excel takes this from the first workbook opened in a session. Close other workbooks first if you see a circular-reference warning.)",
        "2) Press F9 a few times until the 'Max speed' monitor on 'Inputs' stops changing.",
        "",
        "SHEETS:",
        "Inputs   - free-stream speed, streamline spacing, monitors, Reset switch.",
        "Mask     - 1 = solid, 0 = fluid. Default is a semicircular bump of radius 0.3 m. Type 1s to draw your own obstacle.",
        "Psi      - stream function. Dark bands are streamlines.   u_vel / v_vel / Speed - velocity field.",
        "",
        "RULES FOR THE MASK:",
        "- Obstacles must touch the bottom wall (the wall is the psi = 0 streamline; a floating body would need a different constant).",
        "- Keep the inlet and outlet columns fluid.",
        "",
        "SANITY CHECK: set every mask cell to 0 -> psi = U*y, u = U, v = 0 everywhere (uniform stream).",
        "Limitation: ideal flow has no viscosity, separation or wake - this is the textbook potential-flow picture.",
    ])
    wb.move_sheet("Inputs", offset=-(len(wb.sheetnames) - 2))
    wb.save(path)


# --------------------------------------------------------------------------
# 3. lid-driven cavity (stream function - vorticity)
# --------------------------------------------------------------------------
GHIA_RE100 = [  # y, u along vertical centreline x=0.5  (Ghia, Ghia & Shin 1982, Table I, Re=100)
    (1.0000, 1.00000), (0.9766, 0.84123), (0.9688, 0.78871), (0.9609, 0.73722),
    (0.9531, 0.68717), (0.8516, 0.23151), (0.7344, 0.00332), (0.6172, -0.13641),
    (0.5000, -0.20581), (0.4531, -0.21090), (0.2813, -0.15662), (0.1719, -0.10150),
    (0.1016, -0.06434), (0.0703, -0.04775), (0.0625, -0.04192), (0.0547, -0.03717),
    (0.0000, 0.00000),
]


def build_cavity(path, N=33):
    NX = NY = N
    wb = new_workbook()
    gr = grid_range(NX, NY)
    rows = [
        ("RUN CONTROL", None, None, None, "section"),
        ("Reset (1 = initialise, 0 = run)", "reset", 0,
         "Set to 1 for fluid at rest, then back to 0 to run.", "input"),
        ("FLOW", None, None, None, "section"),
        ("Lid speed U [m/s]", "U", 1.0, "Top wall moves in +x; cavity is 1 m x 1 m", "input"),
        ("Reynolds number Re = U*L/nu", "Re", 100, "Re <= ~400 is reliable on this grid. Ghia benchmark validation applies at Re = 100", "input"),
        ("Kinematic viscosity nu [m^2/s]", "nu", "=p_U*1/p_Re", "calculated", "calc"),
        ("NUMERICS", None, None, None, "section"),
        ("Time step dt [s]", "dt", 0.01, "Must satisfy all three limits below", "input"),
        ("Upwind blend theta (0 = central, 1 = upwind)", "theta", 0, "Central is more accurate; raise towards 1 if you push Re high and see wiggles", "input"),
        ("Streamline spacing d(psi)", "dpsi", 0.01, "Dark bands on the Psi sheet are streamlines", "input"),
        ("Grid nodes per side N", "N", N, "Fixed by layout; re-run generate_workbooks.py for another N", "calc"),
        ("Grid spacing h [m]", "h", "=1/(p_N-1)", "", "calc"),
        ("Limit: diffusion  h^2/(4 nu)", "limD", "=p_h^2/(4*p_nu)", "dt must be below this", "calc"),
        ("Limit: viscous-lid  4 nu / U^2", "limV", "=4*p_nu/p_U^2", "dt must be below this", "calc"),
        ("Limit: CFL  h / U", "limC", "=p_h/p_U", "dt must be below this", "calc"),
        ("Stability check", None, '=IF(AND(p_dt<=p_limD,p_dt<=p_limV,p_dt<=p_limC),"STABLE","UNSTABLE - reduce dt")', "", "calc"),
        ("MONITORS", None, None, None, "section"),
        ("Kinetic energy (integral of |V|^2/2)", None, f"=0.5*p_h^2*(SUMSQ(u_vel!{gr})+SUMSQ(v_vel!{gr}))", "Plateaus at steady state (~0.042 for Re = 100 on the default 33x33 grid)", "calc"),
        ("Primary vortex strength  min(psi)", None, f"=MIN(Psi!{gr})", "Ghia et al. Re = 100: -0.1034", "calc"),
        ("Max |omega| (excluding walls)", None, f"=MAX(MAX(Omega!{cell(1, 1)}:{cell(NY - 2, NX - 2)}),-MIN(Omega!{cell(1, 1)}:{cell(NY - 2, NX - 2)}))", "", "calc"),
        ("RMS error vs Ghia (Re = 100)", None, "=Validation!K3", "Only meaningful at Re = 100 and steady state", "calc"),
    ]
    inp = make_inputs_sheet(wb, "2D lid-driven cavity  -  inputs & monitors", rows)

    ps = make_field_sheet(wb, "Psi", NX, NY, "0.000", "Stream function psi   (laplacian psi = -omega;  u = dpsi/dy, v = -dpsi/dx)")
    om = make_field_sheet(wb, "Omega", NX, NY, "0.0", "Vorticity omega   (transport equation, explicit in time)")
    uv = make_field_sheet(wb, "u_vel", NX, NY, "0.00", "Horizontal velocity u")
    vv = make_field_sheet(wb, "v_vel", NX, NY, "0.00", "Vertical velocity v")

    for i in range(NY):
        for j in range(NX):
            a = cell(i, j)
            interior = 0 < i < NY - 1 and 0 < j < NX - 1
            E, W, Nn, S = (cell(i, min(j + 1, NX - 1)), cell(i, max(j - 1, 0)),
                           cell(max(i - 1, 0), j), cell(min(i + 1, NY - 1), j))
            if interior:
                ps[a] = f"=IF(p_reset=1,0,0.25*({E}+{W}+{Nn}+{S}+p_h^2*Omega!{a}))"
                w = a
                om[a] = (
                    f"=IF(p_reset=1,0,{w}+p_dt*("
                    f"-u_vel!{a}*((1-p_theta)*({E}-{W})/(2*p_h)+p_theta*IF(u_vel!{a}>0,({w}-{W}),({E}-{w}))/p_h)"
                    f"-v_vel!{a}*((1-p_theta)*({Nn}-{S})/(2*p_h)+p_theta*IF(v_vel!{a}>0,({w}-{S}),({Nn}-{w}))/p_h)"
                    f"+p_nu*({E}+{W}+{Nn}+{S}-4*{w})/p_h^2))"
                )
                uv[a] = f"=(Psi!{Nn}-Psi!{S})/(2*p_h)"
                vv[a] = f"=-(Psi!{E}-Psi!{W})/(2*p_h)"
            else:
                ps[a] = 0
                # Thom's wall-vorticity formula (corner nodes are never used by the interior stencil)
                corner = i in (0, NY - 1) and j in (0, NX - 1)
                if corner:
                    om[a] = 0
                elif i == 0:
                    om[a] = f"=-2*Psi!{cell(1, j)}/p_h^2-2*p_U/p_h"
                elif i == NY - 1:
                    om[a] = f"=-2*Psi!{cell(NY - 2, j)}/p_h^2"
                elif j == 0:
                    om[a] = f"=-2*Psi!{cell(i, 1)}/p_h^2"
                else:
                    om[a] = f"=-2*Psi!{cell(i, NX - 2)}/p_h^2"
                uv[a] = "=p_U" if i == 0 else 0
                vv[a] = 0
    add_streamline_bands(ps, NX, NY)
    add_colorscale(ps, gr, diverging=True)
    add_colorscale(om, gr, diverging=True)
    add_colorscale(uv, gr, diverging=True)
    add_colorscale(vv, gr, diverging=True)

    # ---- validation vs Ghia et al. 1982
    v = wb.create_sheet("Validation")
    v["A1"] = "Vertical centre line (x = 0.5): u-velocity, spreadsheet vs Ghia, Ghia & Shin (1982), Re = 100"
    v["A1"].font = F_BOLD
    for k, t in enumerate(["y", "u (spreadsheet)"]):
        c = v.cell(2, 1 + k, t)
        c.font = F_BOLD
        c.fill = FILL_HDR
    mid = NX // 2
    for i in range(NY):
        r = 3 + i
        v.cell(r, 1, f"=u_vel!A{GRID_R0 + i}")
        v.cell(r, 2, f"=u_vel!{L(GRID_C0 + mid)}{GRID_R0 + i}")
        for c in (1, 2):
            v.cell(r, c).font = F_BASE
            v.cell(r, c).number_format = "0.000"
    for k, t in enumerate(["y (Ghia)", "u (Ghia)", "row k", "u spreadsheet @ y", "difference"]):
        c = v.cell(2, 4 + k, t)
        c.font = F_BOLD
        c.fill = FILL_HDR
    last = 2 + NY
    A = f"$A$3:$A${last}"
    B = f"$B$3:$B${last}"
    for k, (yy, uu) in enumerate(GHIA_RE100):
        r = 3 + k
        v.cell(r, 4, yy)
        v.cell(r, 5, uu)
        # k = position of the last spreadsheet row with y >= Ghia y (column A is sorted descending)
        v.cell(r, 6, f"=MATCH(D{r}-0.000000001,{A},-1)")
        # linear interpolation between row k and row k+1, with guards at both ends
        v.cell(r, 7, (
            f"=IF(D{r}>=$A$3-0.000000001,$B$3,IF(D{r}<=$A${last}+0.000000001,$B${last},"
            f"INDEX({B},F{r})+(D{r}-INDEX({A},F{r}))*(INDEX({B},F{r}+1)-INDEX({B},F{r}))"
            f"/(INDEX({A},F{r}+1)-INDEX({A},F{r}))))"
        ))
        v.cell(r, 8, f"=G{r}-E{r}")
        v.cell(r, 4).font = F_INPUT
        v.cell(r, 5).font = F_INPUT
        for c in (6, 7, 8):
            v.cell(r, c).font = F_BASE
        for c in (4, 5, 7, 8):
            v.cell(r, c).number_format = "0.0000"
    nG = len(GHIA_RE100)
    v["J2"] = "Summary"
    v["J2"].font = F_BOLD
    v["J3"] = "RMS error vs Ghia"
    v["K3"] = f'=IF(p_Re=100,SQRT(SUMSQ(H3:H{2 + nG})/{nG}),"n/a (Re<>100)")'
    v["K3"].number_format = "0.0000"
    v["J3"].font = F_BASE
    v["K3"].font = F_BASE
    v["D" + str(4 + nG)] = ("Ghia values are typed-in benchmark data (blue). Source: U. Ghia, K.N. Ghia, C.T. Shin, "
                           "J. Comput. Phys. 48 (1982) 387-411, Table I (129x129 grid, Re = 100). "
                           "Differences of a few 1e-2 are expected from the coarser grid here.")
    v["D" + str(4 + nG)].font = Font(name=FONT, size=9, italic=True)
    v.column_dimensions["G"].width = 18
    v.column_dimensions["J"].width = 20
    v.column_dimensions["K"].width = 14
    monitor_chart(v, "M3", "u along vertical centre line (Re = 100)", "u", "y",
                  [(v, 2, 1, 3, last, "Spreadsheet", False),
                   (v, 5, 4, 3, 2 + nG, "Ghia et al. (1982)", True)])

    make_instructions(wb, "2D lid-driven cavity in Excel (stream function - vorticity)", [
        "WHAT THIS SOLVES",
        "Incompressible 2D Navier-Stokes in a 1 m x 1 m cavity; the top lid slides at speed U, the other three walls are still.",
        "  Poisson:    laplacian(psi) = -omega              (Gauss-Seidel sweep each step)",
        "  Transport:  d(omega)/dt + u*d(omega)/dx + v*d(omega)/dy = nu * laplacian(omega)   (explicit Euler)",
        "  Velocity:   u = dpsi/dy,  v = -dpsi/dx",
        "  Walls:      psi = 0;  omega from Thom's formula  omega_wall = -2*psi_inner/h^2 - 2*U_wall/h",
        "No pressure equation is needed - that is what makes this practical in a spreadsheet.",
        "",
        "ONE-TIME SETUP (required):",
        "1) File > Options > Formulas > tick 'Enable iterative calculation'.  Maximum Iterations = 1000, Maximum Change = 0.000000001.",
        "   (Excel takes this from the first workbook opened in a session. Close other workbooks first if you see a circular-reference warning.)",
        "2) Press F9. Each press = ~1000 time steps (dt = 0.01 -> 10 s of flow time). Repeat until 'Kinetic energy' on Inputs stops changing (about 2-3 presses at Re = 100).",
        "",
        "SHEETS:",
        "Inputs - Reset, Re, lid speed, dt, upwind blend, stability limits, monitors.",
        "Psi    - stream function; dark bands are streamlines -> the primary vortex plus the two corner eddies.",
        "Omega / u_vel / v_vel - vorticity and velocity fields (diverging colour scales).",
        "Validation - centre-line u against Ghia et al. (1982) benchmark data at Re = 100, with chart and RMS error.",
        "",
        "TIPS:",
        "- Change Re on Inputs and press F9 again; the field continues from where it is. Use Reset = 1 / 0 for a clean restart.",
        "- Keep dt below all three limits shown on Inputs, otherwise the stability check turns to UNSTABLE and values blow up to error values.",
        "  If that happens: set Reset = 1, lower dt, set Reset = 0.",
        "- Higher Re (> 400) needs a finer grid and/or upwind blending (theta > 0).",
    ])
    wb.move_sheet("Inputs", offset=-(len(wb.sheetnames) - 2))
    wb.save(path)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--out", default=".", help="output directory")
    ap.add_argument("--n-heat", type=int, default=21)
    ap.add_argument("--nx-potential", type=int, default=41)
    ap.add_argument("--ny-potential", type=int, default=21)
    ap.add_argument("--n-cavity", type=int, default=33)
    a = ap.parse_args()
    os.makedirs(a.out, exist_ok=True)
    build_heat(os.path.join(a.out, "1_heat_conduction_2d.xlsx"), a.n_heat)
    build_potential(os.path.join(a.out, "2_potential_flow_2d.xlsx"), a.nx_potential, a.ny_potential)
    build_cavity(os.path.join(a.out, "3_lid_driven_cavity_2d.xlsx"), a.n_cavity)
    print("Workbooks written to", os.path.abspath(a.out))


if __name__ == "__main__":
    main()
